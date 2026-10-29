"""Build stats.json and the panels for the Energize Denver brief.

Inputs: this package's own pulls (code/pull_data.py):
- data/master_<year>_raw.json for 2016-2022 (ArcGIS table attributes,
  exactly as returned: numerics are strings, "None"/"" mean missing)
- data/master_<year>.xlsx for 2023-2025 (Hub Excel items, read here
  with openpyxl; the verifier re-reads the raw CSV dump instead)

Harmonization (the field names change almost every year; matching is
by normalized column name, never by position):
- building id: Denver_Building_ID / Denver_Building_Id / Building_Id /
  Building_ID / "Building ID"
- square footage: Master_Square_Footage / Master_Sq_Ft / "Master Sq Ft"
- site EUI: Site_EUI__kBtu_sq_ft_ / Site_Eui__Kbtu_Sq_Ft_ /
  "Site EUI (kBtu/sq ft)" (kBtu per square foot, as reported)
- weather-normalized site EUI, site energy use (kBtu), primary
  property type (EPA calculated), master property type, compliance
  status, year built: matched by the same normalization.

Definitions and traps, stated up front:
- A numeric value is KNOWN only if it parses as a number and is
  greater than zero. The strings "None", "", "nan", "n/a", zero, and
  negative values map to unknown and are excluded from every EUI
  denominator; the unknown share is printed alongside every rate.
- "Reported energy data" in this package means a row with a known
  site EUI. That is an operational definition, not the file's
  Compliance Status label: the label vocabulary changes over the
  window ("Out of Compliance" through 2019, "Not Submitted" from
  2020, "Pending Revisions" added 2021, lowercase in the Excel
  years), and the cross-tab of known-EUI by label is published in
  stats.json so the two can be compared year by year.
- Years are dataset (energy) years. The internal Reporting Year
  field reads one year later in every file that carries it.
- Building type is the file's Primary Property Type (EPA calculated)
  field, grouped case-insensitively. NOTE (verified in the audit):
  the 2018 file's primary-type field carries assessor-style values
  (Apartment, Condominium) where the other years carry EPA values
  (Multifamily Housing); 2018 is therefore excluded from by-type
  comparisons and the fact is printed wherever types are compared.
- Matched panels join on Denver Building ID, which is unique within
  every year (0 duplicates in all 10 files, asserted in the audit).
"""
import json, pathlib, re
from collections import Counter, defaultdict

import numpy as np
import openpyxl
import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent.parent
YEARS = list(range(2016, 2026))
MISSING = {"", "none", "nan", "n/a", "na", "null"}


def norm(name):
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def num(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return f if f == f else None
    s = str(v).strip().replace(",", "")
    if s.lower() in MISSING:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def pick(row, *prefixes, exact=()):
    for k, v in row.items():
        nk = norm(k)
        if nk in exact or any(nk.startswith(p) for p in prefixes):
            return v
    return None


def load_year(year):
    if year <= 2022:
        raw = json.loads((ROOT / "data" / f"master_{year}_raw.json").read_text())
        rows = [dict(r) for r in raw]
    else:
        wb = openpyxl.load_workbook(ROOT / "data" / f"master_{year}.xlsx",
                                    read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        rows = [dict(zip(header, r)) for r in it]
    out = []
    for r in rows:
        bid = num(pick(r, exact=("denverbuildingid", "buildingid")))
        eui = num(pick(r, "siteeui"))
        wn = num(pick(r, "weathernormalizedsiteeui"))
        energy = num(pick(r, "siteenergyuse"))
        sqft = num(pick(r, exact=("mastersquarefootage", "mastersqft")))
        ptype = pick(r, "primarypropertytype")
        out.append({
            "year": year,
            "building_id": int(bid) if bid is not None else None,
            "compliance_status": str(pick(r, exact=("compliancestatus",)) or "").strip(),
            "property_name": str(pick(r, exact=("propertyname",)) or "").strip(),
            "primary_type_raw": "" if ptype is None else str(ptype).strip(),
            "master_type_raw": str(pick(r, exact=("masterpropertytype",)) or "").strip(),
            "sqft": sqft,
            "site_eui": eui,
            "wn_site_eui": wn,
            "site_energy_kbtu": energy,
            "year_built": num(pick(r, exact=("yearbuilt",))),
            "reporting_year_raw": pick(r, "reportedforyear", "reportingyear"),
        })
    return out


def q(s, p):
    return round(float(np.percentile(s, p)), 1) if len(s) else None


def eui_block(vals):
    s = pd.Series([v for v in vals if v is not None and v > 0], dtype=float)
    return {"n": int(len(s)), "median": q(s, 50), "mean": round(float(s.mean()), 1) if len(s) else None,
            "p25": q(s, 25), "p75": q(s, 75), "min": round(float(s.min()), 1) if len(s) else None,
            "max": round(float(s.max()), 1) if len(s) else None}


def main():
    frames = {y: load_year(y) for y in YEARS}
    panel = pd.DataFrame([r for y in YEARS for r in frames[y]])
    # canonical type display: most common raw casing for each casefold key
    casing = defaultdict(Counter)
    for t in panel["primary_type_raw"]:
        if t:
            casing[t.casefold()][t] += 1
    display = {k: c.most_common(1)[0][0] for k, c in casing.items()}
    panel["primary_type"] = panel["primary_type_raw"].map(
        lambda t: display.get(t.casefold(), "Not reported")
        if t and t != "0" else "Not reported")
    panel["eui_known"] = panel["site_eui"].notna() & (panel["site_eui"] > 0)
    panel["sqft_known"] = panel["sqft"].notna() & (panel["sqft"] > 0)
    panel.to_csv(ROOT / "data" / "building_panel.csv", index=False)

    stats = {"source_pull_date": "2026-10-29",
             "years": [str(y) for y in YEARS],
             "year_convention": "Dataset (energy) year; the internal Reporting Year field reads one year later in every file that carries it (2022 table: 2023; 2025 workbook: 2026).",
             "eui_definition": "Site EUI (kBtu per square foot) as reported in each year's Final Master Dataset. Known means parses as a number and is greater than zero; 'None', blank, zero, and negative values are unknown and excluded from every EUI denominator.",
             "coverage_definition": "Energize Denver benchmarking covers buildings of 25,000 square feet and larger (program description in each dataset). The 2016 file predates the full phase-in and holds fewer buildings.",
             "per_year": {}, "panels": {}, "by_type": {}}

    for y in YEARS:
        d = panel[panel["year"] == y]
        valid = d[d["eui_known"]]
        both = valid[valid["sqft_known"] & valid["site_energy_kbtu"].notna()
                     & (valid["site_energy_kbtu"] > 0)]
        stats["per_year"][str(y)] = {
            "rows": int(len(d)),
            "unique_building_ids": int(d["building_id"].nunique()),
            "duplicate_ids": int(len(d) - d["building_id"].nunique()),
            "eui_known_n": int(len(valid)),
            "eui_known_share_of_rows_pct": round(100 * len(valid) / len(d), 1),
            "eui_zero_n": int(((d["site_eui"] == 0).sum())),
            "eui_missing_n": int((d["site_eui"].isna().sum())),
            "site_eui": eui_block(d["site_eui"]),
            "wn_site_eui_median": eui_block(d["wn_site_eui"])["median"],
            "sqft_total_all_rows": int(d.loc[d["sqft_known"], "sqft"].sum()),
            "sqft_total_eui_known": int(valid.loc[valid["sqft_known"], "sqft"].sum()),
            "sqft_weighted_site_eui": (round(float(both["site_energy_kbtu"].sum()
                                                   / both["sqft"].sum()), 1) if len(both) else None),
            "sqft_weighted_n": int(len(both)),
            "eui_over_300_n": int(((d["site_eui"] > 300)).sum()),
            "eui_over_1000_n": int(((d["site_eui"] > 1000)).sum()),
            "median_excl_eui_over_300": eui_block(
                d.loc[d["site_eui"] <= 300, "site_eui"])["median"],
            "mean_excl_eui_over_300": eui_block(
                d.loc[d["site_eui"] <= 300, "site_eui"])["mean"],
            "compliance_counter": dict(Counter(d["compliance_status"]).most_common()),
            "eui_known_by_compliance": dict(Counter(
                valid["compliance_status"].str.casefold()).most_common()),
            "primary_type_counter_top": dict(Counter(
                valid["primary_type"]).most_common(8)),
        }

    def panel_pair(y0, y1):
        a = panel[(panel["year"] == y0) & panel["eui_known"]].set_index("building_id")
        b = panel[(panel["year"] == y1) & panel["eui_known"]].set_index("building_id")
        ids = a.index.intersection(b.index)
        ea, eb = a.loc[ids, "site_eui"], b.loc[ids, "site_eui"]
        chg = eb - ea
        pct = 100 * chg / ea
        return {"n_buildings": int(len(ids)),
                f"median_eui_{y0}": round(float(ea.median()), 1),
                f"median_eui_{y1}": round(float(eb.median()), 1),
                "median_change": round(float(eb.median() - ea.median()), 1),
                "median_change_pct_of_medians": round(float(100 * (eb.median() - ea.median()) / ea.median()), 1),
                f"mean_eui_{y0}": round(float(ea.mean()), 1),
                f"mean_eui_{y1}": round(float(eb.mean()), 1),
                "median_building_change": round(float(chg.median()), 1),
                "median_building_pct_change": round(float(pct.median()), 1),
                "share_buildings_eui_fell_pct": round(float(100 * (chg < 0).mean()), 1),
                "share_buildings_eui_rose_pct": round(float(100 * (chg > 0).mean()), 1)}

    stats["panels"]["2016_2025"] = panel_pair(2016, 2025)
    stats["panels"]["2019_2025"] = panel_pair(2019, 2025)
    stats["panels"]["2017_2025"] = panel_pair(2017, 2025)
    stats["panels"]["2019_2024"] = panel_pair(2019, 2024)
    every = panel[panel["eui_known"]].groupby("building_id")["year"].nunique()
    stats["panels"]["in_all_10_years_with_known_eui"] = int((every == 10).sum())
    ids25 = set(panel.loc[(panel["year"] == 2025) & panel["eui_known"],
                          "building_id"])
    stats["panels"]["in_2025_and_at_least_5_years"] = int(
        sum(1 for i in ids25 if every.get(i, 0) >= 5))

    def panel_pair_wn(y0, y1):
        a = panel[(panel["year"] == y0) & panel["wn_site_eui"].notna()
                  & (panel["wn_site_eui"] > 0)].set_index("building_id")
        b = panel[(panel["year"] == y1) & panel["wn_site_eui"].notna()
                  & (panel["wn_site_eui"] > 0)].set_index("building_id")
        ids = a.index.intersection(b.index)
        return {"n_buildings": int(len(ids)),
                f"median_wn_eui_{y0}": round(float(a.loc[ids, "wn_site_eui"].median()), 1),
                f"median_wn_eui_{y1}": round(float(b.loc[ids, "wn_site_eui"].median()), 1)}

    stats["panels"]["wn_2016_2025"] = panel_pair_wn(2016, 2025)
    stats["panels"]["wn_2019_2025"] = panel_pair_wn(2019, 2025)

    for y in (2016, 2019, 2025):
        d = panel[(panel["year"] == y) & panel["eui_known"]]
        rows = {}
        for t, g in d.groupby("primary_type"):
            blk = eui_block(g["site_eui"])
            blk["share_of_known_pct"] = round(100 * len(g) / len(d), 1)
            blk["small_count_flag"] = bool(len(g) < 30)
            rows[t] = blk
        stats["by_type"][str(y)] = dict(
            sorted(rows.items(), key=lambda kv: -kv[1]["n"]))

    d25 = panel[(panel["year"] == 2025) & panel["eui_known"]].copy()
    top_eui = d25.sort_values("site_eui", ascending=False).head(10)
    stats["top_buildings_2025_by_eui"] = [
        {"building_id": int(r.building_id), "name": r.property_name,
         "type": r.primary_type, "sqft": (int(r.sqft) if pd.notna(r.sqft) else None),
         "site_eui": float(r.site_eui)} for r in top_eui.itertuples()]
    top_energy = d25[d25["site_energy_kbtu"].notna()].sort_values(
        "site_energy_kbtu", ascending=False).head(5)
    stats["top_buildings_2025_by_energy"] = [
        {"building_id": int(r.building_id), "name": r.property_name,
         "type": r.primary_type, "sqft": (int(r.sqft) if pd.notna(r.sqft) else None),
         "site_eui": float(r.site_eui),
         "site_energy_kbtu": float(r.site_energy_kbtu)} for r in top_energy.itertuples()]

    # Internal control: reported EUI vs energy / sq ft recomputation
    both = panel[panel["eui_known"] & panel["sqft_known"]
                 & panel["site_energy_kbtu"].notna() & (panel["site_energy_kbtu"] > 0)].copy()
    both["recomputed"] = both["site_energy_kbtu"] / both["sqft"]
    agree = (both["recomputed"] - both["site_eui"]).abs() <= np.maximum(1.0, 0.02 * both["site_eui"])
    stats["control_eui_recompute"] = {
        "n_checked": int(len(both)), "n_agree_within_1_or_2pct": int(agree.sum()),
        "agree_pct": round(100 * float(agree.mean()), 1),
        "note": "Site energy use (kBtu) divided by square footage, recomputed per building-year and compared with the reported site EUI. Agreement within max(1.0 kBtu/sq ft, 2 percent) counts as agree."}

    stats["distinct_buildings_all_years"] = int(panel["building_id"].nunique())
    pull_log = json.loads((ROOT / "data" / "pull_log.json").read_text())
    stats["parcels_context"] = {
        "parcels_in_city_register": pull_log["parcels_context"]["live_count"],
        "service": pull_log["parcels_context"]["service"],
        "share_of_parcels_ever_benchmarked_pct": round(
            100 * stats["distinct_buildings_all_years"]
            / pull_log["parcels_context"]["live_count"], 1),
        "note": "Scale context only: benchmarking covers buildings of 25,000 sq ft and larger, so the covered set is a small count of parcels holding a large share of floor area. No parcel-level join is performed; a master-dataset row is a building as reported in Portfolio Manager, which does not map one-to-one onto parcels."}
    stats["type_counter_2018_primary_raw"] = dict(Counter(
        frames[2018][i]["primary_type_raw"] for i in range(len(frames[2018]))).most_common(10))
    (ROOT / "stats.json").write_text(json.dumps(stats, indent=1))
    print(json.dumps({y: stats["per_year"][y]["site_eui"] for y in stats["per_year"]}, indent=1))
    print(json.dumps(stats["panels"], indent=1))
    print("control", stats["control_eui_recompute"])


if __name__ == "__main__":
    main()
